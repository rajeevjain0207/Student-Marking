"""
Essay grading library.

Pure functions that can be called from either the Streamlit UI (`app.py`) or
the command line (`grader.py --submissions ...`). All file reading goes
through `read_text_from_bytes` so it works equally well with paths on disk
and in-memory uploads from a browser.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import sys
import time
from pathlib import Path
from typing import Any

from openai import OpenAI


SUPPORTED_EXTS = {".pdf", ".docx", ".txt", ".md"}


# ---------- file reading ---------------------------------------------------

def read_text_from_bytes(filename: str, data: bytes) -> str:
    """Extract plain text from raw bytes, dispatching on the filename's extension."""
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)
        return "\n".join(parts)

    if suffix in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")

    raise ValueError(f"Unsupported file type: {suffix}")


def read_text_file(path: Path) -> str:
    """Convenience wrapper for the CLI: read a path from disk."""
    return read_text_from_bytes(path.name, path.read_bytes())


def combine_reference_texts(items: list[tuple[str, str]], max_chars: int) -> str:
    """Join (filename, text) pairs into one labelled block, truncated at max_chars."""
    parts: list[str] = []
    total = 0
    truncated = False
    for name, content in items:
        block = f"=== Reference: {name} ===\n{content.strip()}\n"
        if total + len(block) > max_chars:
            block = block[: max(0, max_chars - total)]
            parts.append(block)
            truncated = True
            break
        parts.append(block)
        total += len(block)
    out = "\n".join(parts)
    if truncated:
        out += f"\n\n[Reference materials truncated at {max_chars:,} characters.]"
    return out


# ---------- prompt + grading -----------------------------------------------

def build_grading_prompt(criteria: dict, reference_text: str, essay_text: str) -> str:
    rubric_lines = []
    for i, c in enumerate(criteria["criteria"], 1):
        rubric_lines.append(
            f"{i}. {c['name']}\n"
            f"   - Full marks ({c['full_marks']}): {c['full_marks_description']}\n"
            f"   - Partial marks ({c['partial_marks']}): {c['partial_marks_description']}\n"
            f"   - No marks (0): criterion not met"
        )
    rubric = "\n\n".join(rubric_lines)

    ref_block = (
        f"REFERENCE MATERIALS (the readings/lectures the student should be drawing on):\n\n{reference_text}"
        if reference_text.strip()
        else "REFERENCE MATERIALS: (none provided — judge sourcing on internal coherence)"
    )

    return f"""You are an experienced academic grader. Grade the following student essay against the rubric.

ASSIGNMENT: {criteria.get('assignment_name', 'Essay')}
MAXIMUM MARKS: {criteria['max_marks']}

RUBRIC:

{rubric}

{ref_block}

STUDENT ESSAY:

{essay_text}

INSTRUCTIONS:
- For each criterion, award full / partial / zero marks (use the exact numeric values from the rubric).
- Write 1-3 sentences of reasoning per criterion, citing specific evidence from the essay (a quote, phrase, or clear paraphrase).
- Be fair and consistent. Do not be unduly harsh or unduly lenient.
- The total_marks must equal the sum of the per-criterion marks.

Respond with a JSON object in EXACTLY this shape (no extra keys, no prose outside the JSON):

{{
  "criteria": [
    {{"name": "<criterion name verbatim>", "marks": <number>, "reasoning": "<1-3 sentences with evidence>"}}
  ],
  "total_marks": <number>,
  "overall_feedback": "<2-3 sentence summary of strengths and gaps>"
}}
"""


def grade_essay(
    client: OpenAI,
    model: str,
    criteria: dict,
    reference_text: str,
    essay_text: str,
) -> dict:
    """Send the prompt to OpenAI and return the parsed grading dict."""
    prompt = build_grading_prompt(criteria, reference_text, essay_text)
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system",
             "content": "You are a fair and rigorous academic grader. Always respond with valid JSON matching the requested schema."},
            {"role": "user", "content": prompt},
        ],
        response_format={"type": "json_object"},
        temperature=0.2,
    )
    content = response.choices[0].message.content or ""
    try:
        return json.loads(content)
    except json.JSONDecodeError as e:
        raise RuntimeError(f"Model returned invalid JSON: {e}\n--- raw response ---\n{content}")


# ---------- result rows + CSV ---------------------------------------------

def _slugify(name: str, max_len: int = 30) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
    return s[:max_len].rstrip("_")


def column_keys(criteria: dict) -> list[tuple[str, str]]:
    """Return [(marks_col, reasoning_col), ...] one pair per criterion."""
    keys = []
    used: set[str] = set()
    for i, c in enumerate(criteria["criteria"], 1):
        base = c.get("short_name") or _slugify(c["name"])
        if not base or base in used:
            base = f"c{i}"
        used.add(base)
        keys.append((f"{base}_marks", f"{base}_reasoning"))
    return keys


def fieldnames_for(criteria: dict) -> list[str]:
    fields = ["student_file", "student_name"]
    for marks_col, reasoning_col in column_keys(criteria):
        fields.extend([marks_col, reasoning_col])
    fields.extend(["total_marks", "max_marks", "overall_feedback", "status"])
    return fields


def empty_row(filename: str, criteria: dict) -> dict:
    student_name = Path(filename).stem.replace("_", " ").replace("-", " ").title()
    row = {
        "student_file": filename,
        "student_name": student_name,
        "max_marks": criteria["max_marks"],
        "total_marks": "",
        "overall_feedback": "",
        "status": "ok",
    }
    for marks_col, reasoning_col in column_keys(criteria):
        row[marks_col] = ""
        row[reasoning_col] = ""
    return row


def fill_row_from_result(row: dict, result: dict, criteria: dict) -> dict:
    cols = column_keys(criteria)
    per = result.get("criteria", [])
    if len(per) != len(criteria["criteria"]):
        raise ValueError(f"expected {len(criteria['criteria'])} criteria, got {len(per)}")
    for (marks_col, reasoning_col), graded in zip(cols, per):
        row[marks_col] = graded.get("marks", "")
        row[reasoning_col] = graded.get("reasoning", "")
    row["total_marks"] = result.get("total_marks", "")
    row["overall_feedback"] = result.get("overall_feedback", "")
    return row


def results_to_csv_string(results: list[dict], criteria: dict) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=fieldnames_for(criteria))
    writer.writeheader()
    for r in results:
        writer.writerow(r)
    return buf.getvalue()


# ---------- CLI entry point -----------------------------------------------

def _cli_load_reference(reference_dir: Path | None, max_chars: int) -> str:
    if not reference_dir or not reference_dir.exists():
        return ""
    items = []
    for path in sorted(reference_dir.iterdir()):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTS:
            continue
        try:
            items.append((path.name, read_text_file(path)))
        except Exception as e:
            print(f"  Warning: could not read {path.name}: {e}", file=sys.stderr)
    return combine_reference_texts(items, max_chars)


def _cli_main() -> int:
    parser = argparse.ArgumentParser(description="Grade student essays with OpenAI (CLI).")
    parser.add_argument("--submissions", type=Path, required=True)
    parser.add_argument("--reference", type=Path, default=None)
    parser.add_argument("--criteria", type=Path, default=Path("criteria.json"))
    parser.add_argument("--output", type=Path, default=Path("output/grades.csv"))
    parser.add_argument("--model", default="gpt-4o-mini")
    parser.add_argument("--max-reference-chars", type=int, default=80_000)
    args = parser.parse_args()

    try:
        from dotenv import load_dotenv
        load_dotenv()
    except ImportError:
        pass

    if not os.getenv("OPENAI_API_KEY"):
        print("Error: OPENAI_API_KEY not set.", file=sys.stderr)
        return 1
    if not args.criteria.exists():
        print(f"Error: criteria file not found at {args.criteria}", file=sys.stderr)
        return 1

    criteria = json.loads(args.criteria.read_text(encoding="utf-8"))
    print(f"Rubric: {criteria.get('assignment_name', 'Essay')} (max {criteria['max_marks']})")

    reference_text = _cli_load_reference(args.reference, args.max_reference_chars)
    if reference_text:
        print(f"Reference materials: {len(reference_text):,} chars")

    files = sorted(p for p in args.submissions.iterdir()
                   if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS)
    if not files:
        print(f"Error: no supported files in {args.submissions}", file=sys.stderr)
        return 1

    client = OpenAI()
    results = []
    start = time.time()
    for i, path in enumerate(files, 1):
        print(f"[{i}/{len(files)}] {path.name}")
        row = empty_row(path.name, criteria)
        try:
            essay = read_text_file(path).strip()
            if not essay:
                raise ValueError("essay text is empty after extraction")
            result = grade_essay(client, args.model, criteria, reference_text, essay)
            fill_row_from_result(row, result, criteria)
            print(f"    -> {row['total_marks']} / {criteria['max_marks']}")
        except Exception as e:
            row["status"] = f"error: {e}"
            print(f"    ERROR: {e}", file=sys.stderr)
        results.append(row)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(results_to_csv_string(results, criteria), encoding="utf-8")
    elapsed = time.time() - start
    ok = sum(1 for r in results if r["status"] == "ok")
    print(f"\nDone in {elapsed:.1f}s. {ok}/{len(results)} graded. CSV -> {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(_cli_main())
